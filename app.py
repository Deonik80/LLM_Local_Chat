# Copyright (C) 2026 Deonik80 (https://github.com/Deonik80)
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://gnu.org>.

# LM Studio Chat
from __future__ import annotations
import asyncio, base64, csv, html as html_mod, io, json, mimetypes, os, sys, subprocess, time, uuid
# ---------- auto-install недостающих парсеров документов ----------
# Только лёгкие pure-python пакеты: pypdf / python-docx / openpyxl.
# Тяжёлые опциональные (PyAudio и т.п.) сюда не тянем, чтобы не ломать старт.
_AUTO_DEPS = {
    "pypdf": "pypdf",
    "docx": "python-docx",
    "openpyxl": "openpyxl",
}

def _ensure_doc_deps() -> None:
    missing: list[str] = []
    for mod, pip_name in _AUTO_DEPS.items():
        try:
            __import__(mod)
        except ImportError:
            missing.append(pip_name)
    if not missing:
        return
    print(f"[deps] отсутствуют {missing}, устанавливаю: pip install {' '.join(missing)} ...")
    try:
        subprocess.run(
            [sys.executable, "-m", "pip", "install", *missing],
            check=False, timeout=180,
        )
    except Exception as ex:
        print(f"[deps] автоустановка не удалась: {ex}")
        return
    # повторная проверка — что реально встало
    still = []
    for mod, pip_name in _AUTO_DEPS.items():
        try:
            __import__(mod)
        except ImportError:
            if pip_name in missing:
                still.append(pip_name)
    if still:
        print(f"[deps] не удалось установить: {still}. Выполните вручную: pip install {' '.join(still)}")

_ensure_doc_deps()

from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Callable, Optional
import flet as ft  # Version: 0.86.5
import httpx

# 
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None 

def extract_text(p: str) -> str:
    """..."""
    # ...
    if suf == ".pdf":
        if PdfReader is None:
            raise FileError("Библиотека pypdf не установлена. Установите ее командой: pip install pypdf")
        r = PdfReader(str(path))

# ---------- i18n: RU/EN ----------
LANGS = {
"ru": {
    "title": "LM Studio Chat", "send": "Отправить", "stop": "Стоп",
    "new_chat": "Новый чат", "chats": "Чаты", "search": "Поиск чатов...",
    "delete": "Удалить", "rename": "Переименовать", "export": "Экспорт",
    "theme": "Тема", "font": "Шрифт", "check_conn": "Проверить соединение",
    "summarize": "Сжать историю", "tokens": "токенов",
    "input_hint": "Сообщение... Enter — отправить, Shift+Enter — новая строка",
    "reasoning": "Рассуждения модели", "copy": "Копировать", "copied": "Скопировано!",
    "edit": "Изменить", "retry": "Перегенерировать", "variants": "Вариант",
    "model": "Модель", "preset": "Пресет", "settings": "Настройки",
    "sec_model": "Модель", "sec_prompt": "Промпт",
    "preset_name": "Имя пресета", "preset_name_hint": "Мой пресет",
    "to_preset": "В пресет", "del_preset": "Удалить пресет", "clear": "Очистить",
    "load_txt": "Загрузить .txt", "save_txt": "Сохранить .txt",
    "sys_prompt": "System prompt", "attach": "Прикрепить файл",
    "refresh_models": "Обновить модели", "more": "Ещё",
    "m_compress": "Сжать историю", "m_exp_md": "Экспорт Markdown",
    "m_exp_html": "Экспорт HTML", "m_exp_json": "Экспорт JSON",
    "m_theme": "Светлая/тёмная тема", "m_font_up": "Шрифт +", "m_font_down": "Шрифт −",
    "connected": "Подключено", "no_conn": "Нет связи",
    "empty_chat": "Пустой чат", "only_files": "Только вложения",
    "you": "Вы", "me": "Новый чат",
    "dlg_rename": "Переименовать чат", "chat_name": "Название чата",
    "cancel": "Отмена", "save": "Сохранить",
    "p_ordinary": "Обычный", "p_translator": "Переводчик",
    "p_reviewer": "Ревью кода", "p_simple": "Простыми словами",
    "e_enter_name": "Введите название чата", "e_enter_preset": "Введите имя пресета",
    "e_empty_prompt": "System prompt пуст — нечего сохранять",
    "e_builtin_name": "Это имя встроенного пресета — выберите другое",
    "e_builtin_del": "Встроенные пресеты удалить нельзя",
    "e_no_preset": "Такого пользовательского пресета нет",
    "e_no_msgs": "Нет сообщений", "e_few_msgs": "Мало сообщений",
    "e_open_dlg": "Не удалось открыть диалог: {e}",
    "e_read_file": "Не удалось прочитать файл: {e}",
    "e_save_file": "Не удалось сохранить файл: {ex}",
    "e_save_preset": "Не удалось сохранить пресет: {ex}",
    "e_del_preset": "Не удалось удалить пресет: {ex}",
    "load_cancel": "Загрузка отменена", "save_cancel": "Сохранение отменено",
    "prompt_loaded": "System prompt загружен: {n} ({l} симв.)",
    "prompt_saved": "System prompt сохранён: {p}",
    "prompt_cleared": "System prompt очищен",
    "preset_saved": "Пресет «{n}» сохранён", "preset_deleted": "Пресет «{n}» удалён",
    "exported": "Экспортировано в data/ ({f})",
    "stopped": "Остановлено", "assistant": "Ассистент", "ai": "AI",
    "models_n": "Моделей: {n}", "font_n": "Шрифт {v} (применится к новым сообщениям)",
    "summary_of": "Суммируй диалог кратко:\n{t}",
    "summary_hist": "[Саммари истории]\n{c}",
    "no_link": "Нет связи с {u}: {e}", "net_err": "Ошибка сети: {e}",
    "file_not_found": "Файл не найден: {p}", "img_too_big": "Картинка слишком большая",
    "not_image": "Не картинка: {n}", "file_too_big": "Файл > {m} MB",
    "cant_read": "Не удалось прочитать {n}: {e}",
    "pdf_empty": "[PDF без текстового слоя]", "clipped": "\n\n[...обрезано, всего {n} символов...]",
    "lang": "Язык", "fb_helpful": "Полезно", "fb_bad": "Неточно", "fb_harm": "Опасно",
    "tts_speak": "Озвучить", "stt_mic": "Голосовой ввод",
    "mic_listening": "🎤 Говорите… нажмите микрофон для завершения",
    "mic_stop_title": "Остановить запись",
    "generating": "⏳ Модель думает…",
    "typing": "✍️ Модель печатает…",
    "e_no_tts": "Нет TTS: pip install edge-tts",
    "e_no_stt": "Нет SpeechRecognition: pip install SpeechRecognition PyAudio",
    "copy_code": "Копировать код",
    "model_loading": "Загрузка модели {m}…", "model_loaded": "Модель загружена: {m}",
    "e_load_model": "Не удалось загрузить {m}: {e}",
    "e_no_model_on_server": "Модели {m} нет на сервере",
    "dlg_novision_t": "Модель без поддержки изображений",
    "dlg_novision_c": "Сервер ответил 400 — вероятно, модель не принимает картинки. Повторить запрос без изображений (в историю вставится пометка)?",
    "retry_noimg": "Повторить без картинок",
    "vision_off": "🚫 vision выкл — отправлено без изображений",
    "vision_auto": "🚫 авто: модель без vision — отправлено без изображений",
    "vision_send": "Vision: отправлять изображения",
    "dlg_ctxfull_t": "Контекст переполнен",    "dlg_ctxfull_c": "~{cur} / {cl} токенов. Сжать историю (summarize), продолжить как есть или отменить?",
    "ctx_still_over": "всё ещё переполнено",
    "continue_btn": "Продолжить",
    "m_find": "Найти в чате", "m_rated": "Только оценённые ★", "m_exp_fb": "Экспорт фидбека JSONL",
    "sec_profiles": "Профили связок", "profile": "Профиль",
    "profile_name": "Имя профиля", "profile_name_hint": "Код-ревью Qwen",
    "to_profile": "В профиль", "del_profile": "Удалить профиль",
    "profile_applied": "Профиль «{n}» применён", "profile_saved": "Профиль «{n}» сохранён",
    "profile_deleted": "Профиль «{n}» удалён",
    "sec_env": "Окружение", "env_none": "Окружение не загружено",
    "env_req": "requirements: {n} симв.", "env_vars_n": ".env: {n} vars",
    "req_loaded": "requirements загружен: {n}", "env_loaded": ".env загружен: {n} vars",
},
"en": {
    "title": "LM Studio Chat", "send": "Send", "stop": "Stop",
    "new_chat": "New chat", "chats": "Chats", "search": "Search chats...",
    "delete": "Delete", "rename": "Rename", "export": "Export",
    "theme": "Theme", "font": "Font", "check_conn": "Check connection",
    "summarize": "Summarize history", "tokens": "tokens",
    "input_hint": "Message... Enter to send, Shift+Enter for new line",
    "reasoning": "Model reasoning", "copy": "Copy", "copied": "Copied!",
    "edit": "Edit", "retry": "Regenerate", "variants": "Variant",
    "model": "Model", "preset": "Preset", "settings": "Settings",
    "sec_model": "Model", "sec_prompt": "Prompt",
    "preset_name": "Preset name", "preset_name_hint": "My preset",
    "to_preset": "Save preset", "del_preset": "Delete preset", "clear": "Clear",
    "load_txt": "Load .txt", "save_txt": "Save .txt",
    "sys_prompt": "System prompt", "attach": "Attach file",
    "refresh_models": "Refresh models", "more": "More",
    "m_compress": "Summarize history", "m_exp_md": "Export Markdown",
    "m_exp_html": "Export HTML", "m_exp_json": "Export JSON",
    "m_theme": "Light/dark theme", "m_font_up": "Font +", "m_font_down": "Font −",
    "connected": "Connected", "no_conn": "No connection",
    "empty_chat": "Empty chat", "only_files": "Attachments only",
    "you": "You", "me": "New chat",
    "dlg_rename": "Rename chat", "chat_name": "Chat name",
    "cancel": "Cancel", "save": "Save",
    "p_ordinary": "Default", "p_translator": "Translator",
    "p_reviewer": "Code reviewer", "p_simple": "Simply explained",
    "e_enter_name": "Enter chat name", "e_enter_preset": "Enter preset name",
    "e_empty_prompt": "System prompt is empty — nothing to save",
    "e_builtin_name": "This is a built-in preset name — pick another",
    "e_builtin_del": "Built-in presets cannot be deleted",
    "e_no_preset": "No such custom preset",
    "e_no_msgs": "No messages", "e_few_msgs": "Too few messages",
    "e_open_dlg": "Could not open dialog: {e}",
    "e_read_file": "Could not read file: {e}",
    "e_save_file": "Could not save file: {ex}",
    "e_save_preset": "Could not save preset: {ex}",
    "e_del_preset": "Could not delete preset: {ex}",
    "load_cancel": "Load cancelled", "save_cancel": "Save cancelled",
    "prompt_loaded": "System prompt loaded: {n} ({l} chars)",
    "prompt_saved": "System prompt saved: {p}",
    "prompt_cleared": "System prompt cleared",
    "preset_saved": "Preset “{n}” saved", "preset_deleted": "Preset “{n}” deleted",
    "exported": "Exported to data/ ({f})",
    "stopped": "Stopped", "assistant": "Assistant", "ai": "AI",
    "models_n": "Models: {n}", "font_n": "Font {v} (applies to new messages)",
    "summary_of": "Briefly summarize the dialogue:\n{t}",
    "summary_hist": "[History summary]\n{c}",
    "no_link": "No connection to {u}: {e}", "net_err": "Network error: {e}",
    "file_not_found": "File not found: {p}", "img_too_big": "Image too large",
    "not_image": "Not an image: {n}", "file_too_big": "File > {m} MB",
    "cant_read": "Could not read {n}: {e}",
    "pdf_empty": "[PDF has no text layer]", "clipped": "\n\n[...clipped, {n} chars total...]",
    "lang": "Language", "fb_helpful": "Helpful", "fb_bad": "Inaccurate", "fb_harm": "Harmful",
    "tts_speak": "Speak aloud", "stt_mic": "Voice input",
    "mic_listening": "🎤 Listening… press mic to finish",
    "mic_stop_title": "Stop recording",
    "generating": "⏳ Thinking…",
    "typing": "✍️ Typing…",
    "e_no_tts": "No TTS engine: pip install edge-tts",
    "e_no_stt": "No SpeechRecognition: pip install SpeechRecognition PyAudio",
    "copy_code": "Copy code",
    "model_loading": "Loading model {m}…", "model_loaded": "Model loaded: {m}",
    "e_load_model": "Failed to load {m}: {e}",
    "e_no_model_on_server": "Model {m} not found on server",
    "dlg_novision_t": "Model without image support",
    "dlg_novision_c": "Server replied 400 — the model likely rejects images. Retry without images (a note will be inserted)?",
    "retry_noimg": "Retry without images",
    "vision_off": "🚫 vision off — sent without images",
    "vision_auto": "🚫 auto: no-vision model — sent without images",
    "vision_send": "Vision: send images",
    "dlg_ctxfull_t": "Context full",
    "dlg_ctxfull_c": "~{cur} / {cl} tokens. Summarize history, continue as-is, or cancel?",
    "ctx_still_over": "still over the limit",
    "continue_btn": "Continue",
    "m_find": "Find in chat", "m_rated": "Rated only ★", "m_exp_fb": "Export feedback JSONL",
    "sec_profiles": "Combo profiles", "profile": "Profile",
    "profile_name": "Profile name", "profile_name_hint": "Code-review Qwen",
    "to_profile": "Save profile", "del_profile": "Delete profile",
    "profile_applied": "Profile “{n}” applied", "profile_saved": "Profile “{n}” saved",
    "profile_deleted": "Profile “{n}” deleted",
    "sec_env": "Environment", "env_none": "No environment loaded",
    "env_req": "requirements: {n} chars", "env_vars_n": ".env: {n} vars",
    "req_loaded": "requirements loaded: {n}", "env_loaded": ".env loaded: {n} vars",
},
}
CUR = {"lang": "ru"}
try:
    from localization import LocalizationManager
    from logging_config import setup_logging, get_logger
    _i18n = LocalizationManager(LANGS, default="ru", preset_i18n={
        "Обычный": "p_ordinary", "Переводчик": "p_translator",
        "Ревью кода": "p_reviewer", "Простыми словами": "p_simple"})
    _log = get_logger("app")
except ImportError:  # standalone run without new modules: fall back to legacy
    LocalizationManager = None  # type: ignore
    _i18n = None  # type: ignore
    def setup_logging(*a, **k):  # type: ignore
        import logging as _l
        return _l.getLogger("app")
    def get_logger(name="app"):  # type: ignore
        import logging as _l
        return _l.getLogger(name)
    _log = get_logger("app")
# ---------- 6: presentation-слой (см. views.py; fallback — локальная логика) ----------
try:
    from views import (visible_messages as _visible_messages,  # type: ignore
                       format_feedback as _format_feedback,
                       token_stats as _token_stats,
                       ViewBinder as _ViewBinder)
except ImportError:
    _visible_messages = _format_feedback = _token_stats = _ViewBinder = None  # type: ignore
def tr(key: str, **kw) -> str:
    if _i18n is not None:
        return _i18n.tr(key, **kw)
    s = LANGS.get(CUR["lang"], LANGS["ru"]).get(key, LANGS["ru"].get(key, key))
    try: return s.format(**kw) if kw else s
    except Exception: return s
STR = LANGS["ru"]  # legacy alias
PRESET_I18N = {"Обычный": "p_ordinary", "Переводчик": "p_translator",
               "Ревью кода": "p_reviewer", "Простыми словами": "p_simple"}
def preset_display(key: str) -> str:
    k = PRESET_I18N.get(key)
    return tr(k) if k else key
BASE_URL = os.getenv("LM_STUDIO_URL", "http://localhost:1234/v1")
CHAT_URL = f"{BASE_URL}/chat/completions"
MODEL_URL = f"{BASE_URL}/models"
DEFAULT_MODEL = os.getenv("DEFAULT_MODEL", "default-model")
TIMEOUT = float(os.getenv("REQUEST_TIMEOUT", "180"))
MAX_IMG = int(os.getenv("MAX_IMAGE_MB", "10")) * 1024 * 1024
MAX_TXT = int(os.getenv("MAX_TEXT_MB", "5")) * 1024 * 1024
MAX_CTX = int(os.getenv("MAX_CONTEXT_MESSAGES", "20"))
DATA = Path(__file__).parent / "data"
CHATS = DATA / "chats"; ATTACH = DATA / "attachments"
INDEX_F = DATA / "index.json"; SET_F = DATA / "settings.json"; PRESETS_F = DATA / "presets.json"
PROFILES_F = DATA / "profiles.json"
for p in (CHATS, ATTACH): p.mkdir(parents=True, exist_ok=True)

def _all_presets() -> dict:  # встроенные (на текущем языке) + пользовательские из presets.json
    d = {k: (v.get(CUR["lang"], v["ru"]) if isinstance(v, dict) else v)
         for k, v in BUILTIN_PRESETS.items()}
    try:
        if PRESETS_F.is_file():
            raw = json.loads(PRESETS_F.read_text("utf-8"))
            if isinstance(raw, dict):
                d.update({str(k): str(v) for k, v in raw.items()})
    except Exception: pass
    return d
def _save_custom_presets(custom: dict):
    PRESETS_F.write_text(json.dumps(custom, ensure_ascii=False, indent=2), "utf-8")
def _load_custom_presets() -> dict:
    try:
        if PRESETS_F.is_file():
            raw = json.loads(PRESETS_F.read_text("utf-8"))
            if isinstance(raw, dict): return {str(k): str(v) for k, v in raw.items()}
    except Exception: pass
    return {}

DEFAULT_SETTINGS = {"system_prompt": "", "temperature": 0.7, "top_p": 1.0,
    "repeat_penalty": 1.0, "seed": -1, "max_tokens": 2048,
    "context_length": int(os.getenv("CONTEXT_LENGTH", "8192")), "send_images": True,
    "model": os.getenv("DEFAULT_MODEL", ""), "no_vision_models": [],
    "window_width": 1100, "window_height": 860, "window_left": None, "window_top": None,
    "window_maximized": False,
    "lang": "ru",
    "theme": "dark", "font_scale": 1.0, "preset": "Обычный"}
BUILTIN_PRESETS = {
    "Обычный": {"ru": "", "en": ""},
    "Переводчик": {"ru": "Ты профессиональный переводчик. Переводи точно, сохраняй стиль.",
                   "en": "You are a professional translator. Translate accurately and keep the style."},
    "Ревью кода": {"ru": "Ты senior-разработчик. Находи баги, предлагай исправления с кодом.",
                   "en": "You are a senior developer. Find bugs and suggest fixes with code."},
    "Простыми словами": {"ru": "Объясняй просто, с примерами и аналогиями.",
                         "en": "Explain simply, with examples and analogies."}}
PRESETS = _all_presets()
# (1) Дизайн-система: палитра + метрики в одном месте
S = {"radius": 12, "bubble_radius": 14, "pad": 12, "gap": 8,
     "sidebar_w": 260, "input_max": 4000, "bubble_margin": 110}
THEMES = {"dark": {"bg": "#1E1E1E", "panel": "#252526", "border": "#333333",
    "ubg": "#1976D2", "abg": "#2D2D2D", "utc": "#FFFFFF", "atc": "#E0E0E0",
    "muted": "#9A9A9A", "accent": "#4FA3FF", "hover": "#3A3A3C",
    "input_bg": "#252526", "chip_bg": "#333333", "shadow": "#000000"},
    "light": {"bg": "#F2F3F7", "panel": "#FFFFFF", "border": "#E0E3EB",
    "ubg": "#1976D2", "abg": "#FFFFFF", "utc": "#FFFFFF", "atc": "#222222",
    "muted": "#6B7280", "accent": "#1976D2", "hover": "#EEF1F6",
    "input_bg": "#FFFFFF", "chip_bg": "#EEF1F6", "shadow": "#B0B5C0"}}

# ---------- модели (Pydantic v2, см. models.py; fallback — dataclasses) ----------
try:
    from models import Attachment, ChatMessage  # type: ignore
except ImportError:
    @dataclass
    class Attachment:
        path: str; mime: Optional[str] = None; b64: Optional[str] = None
        def to_dict(self): return asdict(self)
        @staticmethod
        def from_dict(d): return Attachment(d["path"], d.get("mime"), d.get("b64"))

    @dataclass
    class ChatMessage:
        text: str; is_user: bool; ts: float = field(default_factory=time.time)
        attachments: list = field(default_factory=list)
        variants: list = field(default_factory=list)  # альтернативные ответы (4)
        rating: Optional[int] = None  # 1..5 feedback (assistant only)
        feedback_type: Optional[str] = None  # helpful|inaccurate|harmful|other
        def to_dict(self):
            return {"text": self.text, "is_user": self.is_user, "ts": self.ts,
                    "attachments": [a.to_dict() if isinstance(a, Attachment) else a for a in self.attachments],
                    "variants": self.variants, "rating": self.rating,
                    "feedback_type": self.feedback_type}
        @staticmethod
        def from_dict(d):
            atts = [Attachment.from_dict(a) if isinstance(a, dict) and "path" in a else a for a in d.get("attachments", [])]
            return ChatMessage(d["text"], d["is_user"], d.get("ts", 0), atts, d.get("variants", []),
                               d.get("rating"), d.get("feedback_type"))

class FileError(Exception): pass
class StreamCancelled(Exception): pass

# ---------- 2: парсинг файлов ----------
def encode_image(p: str):
    path = Path(p)
    if not path.is_file(): raise FileError(tr("file_not_found", p=p))
    if path.stat().st_size > MAX_IMG: raise FileError(tr("img_too_big"))
    mime, _ = mimetypes.guess_type(path.name)
    if not mime or not mime.startswith("image/"): raise FileError(tr("not_image", n=path.name))
    return base64.b64encode(path.read_bytes()).decode(), mime

def extract_text(p: str) -> str:
    """#2: txt/py/md/csv/pdf/docx/xlsx -> текст с обрезкой."""
    path = Path(p)
    if not path.is_file(): raise FileError(tr("file_not_found", p=p))
    suf = path.suffix.lower()
    try:
        if suf == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(path))
            t = "\n".join((pg.extract_text() or "") for pg in r.pages)
            return _clip(t or tr("pdf_empty"))
        if suf == ".docx":
            from docx import Document
            return _clip("\n".join(x.text for x in Document(str(path)).paragraphs))
        if suf in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:5]:
                out.append(f"## {ws.title}")
                for row in list(ws.iter_rows(values_only=True))[:100]:
                    out.append(" | ".join("" if v is None else str(v) for v in row))
            return _clip("\n".join(out))
        if suf == ".csv":
            with open(path, encoding="utf-8-sig") as f:
                rows = list(csv.reader(f))[:200]
            return _clip("\n".join(" | ".join(r) for r in rows))
        # текстовые
        data = path.read_bytes()
        if len(data) > MAX_TXT: raise FileError(tr("file_too_big", m=MAX_TXT // 1048576))
        return _clip(data.decode("utf-8"))
    except FileError: raise
    except Exception as e: raise FileError(tr("cant_read", n=path.name, e=e))

def _clip(t: str, lim: int = 20000) -> str:
    return t if len(t) <= lim else t[:lim] + tr("clipped", n=len(t))

def estimate_tokens(s: str) -> int:  # #3 грубая оценка (канон — api_payload)
    try:
        from api_payload import estimate_tokens as _est  # type: ignore
        return _est(s)
    except ImportError:
        return max(1, len(s or "") // 4)

# ---------- 7: сетевой клиент (см. lm_client.py; fallback — локальный класс) ----------
try:
    from lm_client import LmClient as _LmClient, StreamCancelled  # type: ignore
    def LmClient():  # type: ignore # factory: та же сигнатура вызова, что раньше
        return _LmClient(MODEL_URL, CHAT_URL, base_url=BASE_URL, timeout=TIMEOUT, tr=tr, log=_log)
except ImportError:
    class StreamCancelled(Exception): pass
    class LmClient:
        def __init__(self): self._c = httpx.AsyncClient(timeout=TIMEOUT); self._cancel = False
        def cancel(self): self._cancel = True
        async def close(self): self._cancel = True; await self._c.aclose()
        async def fetch_models(self) -> list[str]:
            last = None
            for i in range(3):  # ретраи
                try:
                    r = await self._c.get(MODEL_URL); r.raise_for_status()
                    return [m["id"] for m in r.json().get("data", []) if "id" in m]
                except Exception as e: last = e; await asyncio.sleep(1 * (i + 1))
            raise RuntimeError(tr("no_link", u=BASE_URL, e=last))
        async def chat_stream(self, msgs, model, s: dict, on_delta: Callable):
            self._cancel = False
            payload = {"model": model, "messages": msgs, "temperature": s["temperature"],
                "top_p": s.get("top_p", 1.0), "max_tokens": s["max_tokens"], "stream": True}
            if s.get("seed", -1) >= 0: payload["seed"] = s["seed"]
            if s.get("repeat_penalty", 1.0) != 1.0: payload["repeat_penalty"] = s["repeat_penalty"]
            content, reasoning = "", ""
            for attempt in range(2):  # автореконнект: 1 ретрай при обрыве SSE
                try:
                    async with self._c.stream("POST", CHAT_URL, json=payload) as r:
                        r.raise_for_status()
                        async for line in r.aiter_lines():
                            if self._cancel: raise StreamCancelled()
                            if not line or not line.strip().startswith("data:"): continue
                            d = line.strip()[6:]
                            if d == "[DONE]": break
                            try: j = json.loads(d)
                            except ValueError: continue
                            ch = j.get("choices", [{}])[0] if isinstance(j.get("choices"), list) else {}
                            delta = ch.get("delta", {}) or {}
                            if delta.get("reasoning_content"):
                                reasoning += delta["reasoning_content"]; on_delta("reasoning", delta["reasoning_content"])
                            if delta.get("content"):
                                content += delta["content"]; on_delta("content", delta["content"])
                    break
                except StreamCancelled: raise
                except httpx.HTTPStatusError as e:
                    body = ""
                    try: body = e.response.text[:300]
                    except Exception: pass
                    raise RuntimeError(f"HTTP {e.response.status_code}: {body or e}")
                except httpx.HTTPError as e:
                    if self._cancel: raise StreamCancelled() from e
                    if content or reasoning or attempt == 1:
                        raise RuntimeError(tr("net_err", e=e))
                    await asyncio.sleep(1.5 * (attempt + 1))  # backoff перед ретраем
                    continue
            return content, reasoning

# ---------- 1: хранилище чатов (см. repositories.py; fallback — локальные функции) ----------
try:
    from repositories import ChatRepository as _ChatRepository, SettingsRepository as _SettingsRepository  # type: ignore
    _chat_repo = _ChatRepository(CHATS, INDEX_F)
    _set_repo = _SettingsRepository(SET_F, DEFAULT_SETTINGS)
    def load_index() -> list: return _chat_repo.load_index()
    def save_index(idx): _chat_repo.save_index(idx)
    def chat_path(cid): return _chat_repo.chat_path(cid)
    def load_chat(cid) -> list: return _chat_repo.load_chat(cid)
    def save_chat(cid, msgs): _chat_repo.save_chat(cid, msgs)
    def load_settings() -> dict: return _set_repo.load()
    def save_settings(s): _set_repo.save(s)
except ImportError:
    def load_index() -> list:
        if INDEX_F.is_file():
            try: return json.loads(INDEX_F.read_text("utf-8"))
            except Exception: pass
        return []
    def save_index(idx): INDEX_F.write_text(json.dumps(idx, ensure_ascii=False, indent=2), "utf-8")
    def chat_path(cid): return CHATS / f"{cid}.json"
    def load_chat(cid) -> list:
        try: return [ChatMessage.from_dict(m) for m in json.loads(chat_path(cid).read_text("utf-8"))]
        except Exception: return []
    def save_chat(cid, msgs): chat_path(cid).write_text(json.dumps([m.to_dict() for m in msgs], ensure_ascii=False, indent=2), "utf-8")
    def load_settings() -> dict:
        s = dict(DEFAULT_SETTINGS)
        if SET_F.is_file():
            try: s.update(json.loads(SET_F.read_text("utf-8")))
            except Exception: pass
        return s
    def save_settings(s): SET_F.write_text(json.dumps(s, ensure_ascii=False, indent=2), "utf-8")

# ---------- 4: маппинг LLM API (см. api_payload.APIPayloadBuilder) ----------
try:
    from api_payload import APIPayloadBuilder as _APIPayloadBuilder  # type: ignore
    _payload = _APIPayloadBuilder(max_ctx_messages=MAX_CTX, extract_text=extract_text)
    def build_api(messages: list[ChatMessage], system: str, context_tokens: int = 0,
                  strip_images: bool = False) -> list[dict]:
        return _payload.build(messages, system, context_tokens, strip_images=strip_images)
    def format_env_block(settings: dict) -> str:
        return _payload.format_env_block(settings)
    def effective_system(settings: dict, system: str) -> str:
        return _payload.effective_system(settings, system)
except ImportError:
    def build_api(messages: list[ChatMessage], system: str, context_tokens: int = 0) -> list[dict]:
        api = []
        if (system or "").strip(): api.append({"role": "system", "content": system.strip()})
        hist = messages[-MAX_CTX:]
        if context_tokens and context_tokens > 0:
            # trim oldest while estimated history tokens exceed the limit (keep at least last 2)
            while len(hist) > 2 and sum(estimate_tokens(m.text or "") for m in hist) > context_tokens:
                hist = hist[1:]
        for m in hist:
            if m.is_user:
                parts = [{"type": "text", "text": m.text}]
                for a in m.attachments:
                    if isinstance(a, Attachment) and a.mime and a.mime.startswith("image/"):
                        if a.b64:
                            parts.append({"type": "image_url", "image_url": {"url": f"data:{a.mime};base64,{a.b64}"}})
                        else:
                            parts.append({"type": "text", "text": f"[image unavailable (file not found): {Path(a.path).name}]"})
                    elif isinstance(a, Attachment):
                        parts.append({"type": "text", "text": f"--- {Path(a.path).name} ---\n{extract_text(a.path)}"})
                api.append({"role": "user", "content": parts})
            else: api.append({"role": "assistant", "content": m.text})
        return api

    def format_env_block(settings: dict) -> str:
        """System Context block from requirements.txt / .env stored in settings."""
        req = (settings.get("env_requirements") or "").strip()
        env_vars = settings.get("env_vars") or {}
        if not req and not env_vars: return ""
        parts = ["[System Context: Environment]"]
        if req: parts.append(f"requirements:\n{req[:8000]}")
        if env_vars:
            parts.append("environment variables:\n" + "\n".join(f"- {k}={v}" for k, v in env_vars.items()))
            parts.append("Model must respect these library versions and env vars.")
        return "\n".join(parts)

    def effective_system(settings: dict, system: str) -> str:
        env_block = format_env_block(settings)
        blocks = [b for b in [env_block, (system or "").strip()] if b]
        return "\n\n".join(blocks)

# ---------- UI ----------
async def main(page: ft.Page):
    setup_logging()
    settings = load_settings()
    th = THEMES[settings.get("theme", "dark")]
    fs = float(settings.get("font_scale", 1.0))
    CUR["lang"] = settings.get("lang", "ru") if settings.get("lang") in LANGS else "ru"
    if _i18n is not None:
        try: _i18n.set_lang(CUR["lang"])
        except ValueError: pass
    _log.info("app started (lang=%s)", CUR["lang"])
    UI = {}  # ссылки на контролы для apply_lang()
    page.title = tr("title"); page.bgcolor = th["bg"]
    try:  # восстановить геометрию окна с прошлого запуска
        if settings.get("window_maximized"):
            page.window.maximized = True
        else:
            page.window.width = max(600, int(settings.get("window_width", 1100) or 1100))
            page.window.height = max(500, int(settings.get("window_height", 860) or 860))
            if settings.get("window_left") is not None:
                page.window.left = int(settings["window_left"])
            if settings.get("window_top") is not None:
                page.window.top = int(settings["window_top"])
    except Exception as ex:
        _log.warning("restore window geometry failed: %s", ex)
    client = LmClient()
    try:
        from chat_store import ChatStore as _ChatStore  # type: ignore
        store = _ChatStore(chat_repo=_chat_repo, settings_repo=_set_repo,
                           payload=_payload, log=_log)
        state = store.state  # единственное состояние; legacy-код работает с тем же объектом
        state.setdefault("loaded_model", None)
        state.setdefault("model_touched", False)
    except (ImportError, NameError):
        store = None
        state = {"cid": None, "msgs": [], "files": [], "sending": False, "stick": True,
                 "conn_ok": False, "conn_custom": None, "chat_filter": "", "rated_only": False,
                 "loaded_model": None, "model_touched": False}

    # --- виджеты ---
    chat_list = ft.Column(spacing=2, scroll=ft.ScrollMode.AUTO, expand=True)
    search = ft.TextField(hint_text=tr("search"), dense=True,
                          prefix_icon=ft.Icons.SEARCH_OUTLINED, border_radius=S["radius"])
    chat_box = ft.Column(spacing=8, scroll=ft.ScrollMode.AUTO, expand=True, auto_scroll=True)
    status = ft.Text("", size=12, color=th["muted"])
    tok_label = ft.Text("", size=11, color=th["muted"])
    ctx_bar = ft.ProgressBar(width=110, value=0, bar_height=6,
                             color=th["accent"], bgcolor=th["border"])
    counter = ft.Text(f"0 / {S['input_max']}", size=11, color=th["muted"])
    err = ft.Container(visible=False, bgcolor="#B71C1C", border_radius=6, padding=10,
                       content=ft.Text("", color="white"))
    inp = ft.TextField(hint_text=tr("input_hint"), multiline=True, min_lines=1, max_lines=5,
                       expand=True, shift_enter=True)
    model_dd = ft.Dropdown(label=tr("model"), value=DEFAULT_MODEL, width=200,
                           options=[ft.DropdownOption(key=DEFAULT_MODEL, text=DEFAULT_MODEL)])
    dot = ft.Container(width=10, height=10, border_radius=5, bgcolor="#777")
    conn_t = ft.Text("—", size=11, color=th["muted"])
    sys_f = ft.TextField(label=tr("sys_prompt"), value=settings.get("system_prompt", ""), multiline=True, min_lines=2, max_lines=3)
    preset_dd = ft.Dropdown(label=tr("preset"), value=settings.get("preset", "Обычный"),
                            options=[ft.DropdownOption(key=k, text=preset_display(k)) for k in PRESETS])
    tmp = ft.Slider(min=0, max=2, divisions=40, value=settings.get("temperature", 0.7), expand=True)
    topp = ft.Slider(min=0.1, max=1.0, divisions=18, value=settings.get("top_p", 1.0), expand=True)
    maxt = ft.TextField(label="max_tokens", value=str(settings.get("max_tokens", 2048)), width=120)
    seed = ft.TextField(label="seed (-1=off)", value=str(settings.get("seed", -1)), width=120)
    ctxlen = ft.TextField(label="Context Length tokens", value=str(settings.get("context_length", 8192)), width=170,
                          hint_text="e.g. 8192")
    send_images_cb = ft.Checkbox(label=tr("vision_send"),
                                 value=bool(settings.get("send_images", True)))
    attach_row = ft.Row(spacing=8, wrap=True)
    clip = ft.Clipboard(); page.services.extend([ft.FilePicker(), clip])
    picker: ft.FilePicker = page.services[0]

    def persist():
        try: cl = int(str(ctxlen.value or 8192))
        except ValueError: cl = 8192
        cl = min(1000000, max(512, cl))
        settings.update(system_prompt=sys_f.value or "", temperature=float(tmp.value or 0.7),
            top_p=float(topp.value or 1.0), max_tokens=int(str(maxt.value or 2048)),
            seed=int(str(seed.value or -1)), preset=preset_dd.value, context_length=cl,
            send_images=bool(send_images_cb.value))
        # модель пишем только после реального выбора пользователем, иначе стартовый
        # open_chat()->persist() затрёт сохранённую модель плейсхолдером до load_models()
        if state.get("model_touched") and (model_dd.value or ""):
            settings["model"] = model_dd.value
        save_settings(settings)
    def show_e(t): err.content = ft.Text(t, color="white"); err.visible = True; page.update()
    def hide_e(): err.visible = False; page.update()
    def set_conn(ok, label=None):
        if store is not None: store.set_conn(ok, label)  # пишет в тот же state + emit
        else: state["conn_ok"] = ok; state["conn_custom"] = label
        dot.bgcolor = "#4CAF50" if ok else "#EF5350"
        conn_t.value = label if label is not None else (tr("connected") if ok else tr("no_conn"))
        page.update()

    # --- (4) сайдбар: активный акцент + превью + счётчик ---
    def chat_preview(cid) -> tuple[str, int]:
        try:
            ms = load_chat(cid)
            if not ms: return (tr("empty_chat"), 0)
            last = next((m.text for m in reversed(ms) if (m.text or "").strip()), "")
            return ((last[:42] + "…") if len(last) > 42 else (last or tr("only_files")), len(ms))
        except Exception: return ("", 0)
    def refresh_sidebar():
        idx = load_index(); q = (search.value or "").lower()
        chat_list.controls.clear()
        for c in idx:
            if q and q not in c["title"].lower(): continue
            cid = c["id"]
            sel = cid == state["cid"]
            prev, n = chat_preview(cid)
            badge = ft.Container(content=ft.Text(str(n), size=10, color="white"),
                bgcolor=th["accent"], border_radius=8, padding=ft.Padding.symmetric(vertical=2, horizontal=6)) if n else ft.Container()
            chat_list.controls.append(ft.Container(
                bgcolor=th["hover"] if sel else None,
                border_radius=S["radius"], padding=8,
                border=ft.Border.all(1, th["accent"]) if sel else ft.Border.all(1, th["border"]),
                on_click=lambda e, x=cid: open_chat(x),
                content=ft.Row([
                    ft.Column([ft.Text(c["title"][:28], size=13,
                                       weight=ft.FontWeight.BOLD if sel else ft.FontWeight.NORMAL,
                                       color=th["atc"]),
                               ft.Text(prev, size=11, color=th["muted"])], spacing=1, expand=True),
                    badge,
                    ft.IconButton(ft.Icons.EDIT_OUTLINED, icon_size=14, tooltip=tr("rename"),
                        on_click=lambda e, x=cid: rename_chat(x)),
                    ft.IconButton(ft.Icons.DELETE_OUTLINE, icon_size=14, tooltip=tr("delete"),
                        on_click=lambda e, x=cid: del_chat(x))],
                    vertical_alignment=ft.CrossAxisAlignment.CENTER)))
        page.update()
    def new_chat():
        cid = uuid.uuid4().hex[:8]
        idx = load_index(); idx.insert(0, {"id": cid, "title": tr("me"), "ts": time.time()})
        save_index(idx); open_chat(cid)
    def open_chat(cid):
        persist()
        state["cid"] = cid; state["msgs"] = load_chat(cid); state["files"] = []
        state["chat_filter"] = ""; state["rated_only"] = False
        chat_box.controls.clear(); attach_row.controls.clear()
        for m in state["msgs"]: add_bubble(m)
        upd_tokens(); refresh_sidebar(); page.update()
    def del_chat(cid):
        try: chat_path(cid).unlink(missing_ok=True)
        except Exception: pass
        save_index([c for c in load_index() if c["id"] != cid])
        if state["cid"] == cid:
            idx = load_index(); open_chat(idx[0]["id"]) if idx else new_chat(); return
        refresh_sidebar()
    def rename_chat(cid):
        idx = load_index()
        cur = next((c.get("title", "") for c in idx if c["id"] == cid), "")
        name_f = ft.TextField(label=tr("chat_name"), value=cur, autofocus=True)
        def do_save(e=None):
            new = (name_f.value or "").strip()[:60]
            if not new: show_e(tr("e_enter_name")); return
            idx2 = load_index()
            for c in idx2:
                if c["id"] == cid: c["title"] = new
            save_index(idx2); close_dlg(); refresh_sidebar()
        def close_dlg(e=None):
            try: page.pop_dialog()
            except Exception: pass
            page.update()
        dlg = ft.AlertDialog(title=ft.Text(tr("dlg_rename")),
            content=name_f, actions=[ft.TextButton(tr("cancel"), on_click=close_dlg),
            ft.TextButton(tr("save"), on_click=do_save)], actions_alignment=ft.MainAxisAlignment.END)
        page.show_dialog(dlg)

    # --- автопрокрутка вниз (со "прилипанием": не дёргаем, если юзер ушёл вверх) ---
    def scroll_end():
        async def _go():
            await asyncio.sleep(0.05)  # дать UI дорисовать новый контрол
            try: await chat_box.scroll_to(offset=999999, duration=200)
            except Exception: pass
        try: page.run_task(_go)
        except Exception:
            try: asyncio.get_running_loop().create_task(_go())
            except Exception: pass
    def on_chat_scroll(e):
        # e.pixels / e.max_scroll_extent могут отсутствовать — аккуратно
        try:
            px = getattr(e, "pixels", None); mx = getattr(e, "max_scroll_extent", None)
            state["stick"] = True if mx is None else (mx - (px or 0) < 120)
        except Exception: state["stick"] = True
    try: chat_box.on_scroll = on_chat_scroll
    except Exception: pass

    # --- #9 пузырь: аватар, время, reasoning-панель, edit/variants ---
    def add_bubble(m: ChatMessage):
        fs2 = fs; T = th
        av = ft.CircleAvatar(content=ft.Text("👤" if m.is_user else "🤖"), radius=14)
        tstr = time.strftime("%H:%M", time.localtime(m.ts or time.time()))
        md = ft.Markdown(m.text or "…", selectable=True, extension_set=ft.MarkdownExtensionSet.GITHUB_WEB,
            code_theme=ft.MarkdownCodeTheme.ATOM_ONE_DARK,
            code_style_sheet=ft.MarkdownStyleSheet(  # стили именно код-блоков (md_style_sheet их не касается)
                code_text_style=ft.TextStyle(color=T["atc"], size=int(13 * fs2)),
                codeblock_padding=8,
                codeblock_decoration=ft.BoxDecoration(
                    bgcolor=T["panel"], border=ft.Border.all(1, T["border"]),
                    border_radius=S["bubble_radius"])),
            md_style_sheet=ft.MarkdownStyleSheet(
                p_text_style=ft.TextStyle(color=T["utc"] if m.is_user else T["atc"], size=int(14 * fs2)),
                a_text_style=ft.TextStyle(color=T["utc"] if m.is_user else T["accent"],
                                          bgcolor="transparent",
                                          size=int(14 * fs2)),
                blockquote_padding=8,
                blockquote_decoration=ft.BoxDecoration(  # дефолт blue.shade100 нечитаем
                    bgcolor=T["panel"], border=ft.Border.all(1, T["border"]),
                    border_radius=S["bubble_radius"])))
        col = ft.Column(spacing=2, controls=[md])
        for a in m.attachments:
            if isinstance(a, Attachment) and a.mime and a.mime.startswith("image/") and a.b64:
                col.controls.insert(0, ft.Image(src=f"data:{a.mime};base64,{a.b64}", width=200, height=150, fit=ft.BoxFit.CONTAIN))
            elif isinstance(a, Attachment):
                col.controls.insert(0, ft.Text(f"📎 {Path(a.path).name}", size=12, color=T["utc"] if m.is_user else T["atc"]))
        # (2) Пузыри: заголовок + ~70% ширины через отступ + граница/тень
        sender = tr("you") if m.is_user else "LM Studio"
        head = ft.Row([ft.Text(sender, size=11, weight=ft.FontWeight.BOLD,
                               color=T["utc"] if m.is_user else T["accent"]),
                       ft.Text(tstr, size=10, color=T["muted"])], spacing=6)
        col = ft.Column(spacing=4, controls=[head] + col.controls)
        bubble = ft.Container(content=col, padding=S["pad"] + 2,
            expand=True,  # занять ширину flex-строки: Markdown получает границу и переносит текст
            border_radius=S["bubble_radius"],
            border=ft.Border.all(1, T["border"]) if not m.is_user else None,
            shadow=ft.BoxShadow(blur_radius=8, spread_radius=0, color=T["shadow"],
                                offset=ft.Offset(0, 2)) if settings.get("theme") == "light" and not m.is_user else None,
            bgcolor=T["ubg"] if m.is_user else T["abg"],
            margin=ft.Margin(left=S["bubble_margin"], right=0, top=0, bottom=0) if m.is_user
                   else ft.Margin(left=0, right=S["bubble_margin"], top=0, bottom=0))
        row = ft.Row([bubble, av] if m.is_user else [av, bubble],
                     alignment=ft.MainAxisAlignment.END if m.is_user else ft.MainAxisAlignment.START,
                     vertical_alignment=ft.CrossAxisAlignment.START)
        wrap = ft.Column(spacing=2, controls=[row])
        # действия
        acts = ft.Row(spacing=0)
        async def copy(e): await clip.set(m.text); status.value = tr("copied"); page.update()
        acts.controls.append(ft.IconButton(ft.Icons.COPY, icon_size=15, tooltip=tr("copy"), on_click=copy))
        async def speak_msg(e):
            try:
                from voice import speak, is_playing, stop_playback  # type: ignore
            except ImportError:
                show_e(tr("e_no_tts")); return
            if is_playing():  # повторный клик — стоп
                stop_playback(); status.value = ""; page.update(); return
            status.value = "🔊…"; page.update()
            try:
                await asyncio.get_running_loop().run_in_executor(
                    None, speak, m.text, CUR["lang"])
                status.value = ""; page.update()
            except Exception as ex: show_e(str(ex))
        acts.controls.append(ft.IconButton(ft.Icons.VOLUME_UP_OUTLINED, icon_size=15,
                                           tooltip=tr("tts_speak"), on_click=speak_msg))
        if m.is_user:
            async def edit(e):
                inp.value = m.text; page.update()  # #4: правка через поле ввода
                state["msgs"].remove(m); chat_box.controls.remove(wrap); save_chat(state["cid"], state["msgs"]); upd_tokens(); page.update()
            async def dele(e):
                state["msgs"].remove(m); chat_box.controls.remove(wrap); save_chat(state["cid"], state["msgs"]); upd_tokens(); page.update()
            acts.controls += [ft.IconButton(ft.Icons.EDIT, icon_size=15, tooltip=tr("edit"), on_click=edit),
                              ft.IconButton(ft.Icons.DELETE, icon_size=15, tooltip=tr("delete"), on_click=dele)]
        else:
            async def regen(e): await regenerate()
            acts.controls.append(ft.IconButton(ft.Icons.REFRESH, icon_size=15, tooltip=tr("retry"), on_click=regen))
            if "```" in (m.text or ""):  # копировать код из markdown-блоков
                async def copy_code(e):
                    parts = (m.text or "").split("```")
                    code = "\n\n".join(p.split("\n", 1)[1] if "\n" in p else p for p in parts[1::2])
                    await clip.set(code or m.text); status.value = tr("copied"); page.update()
                acts.controls.append(ft.IconButton(ft.Icons.CODE_OUTLINED, icon_size=15,
                                                   tooltip=tr("copy_code"), on_click=copy_code))
            fb_label = ft.Text("", size=11, color=T["muted"])
            def refresh_fb():
                if _format_feedback is not None:
                    fb_label.value = _format_feedback(m)
                else:
                    bits = []
                    if m.rating: bits.append(f"★{m.rating}")
                    if m.feedback_type: bits.append(str(m.feedback_type))
                    fb_label.value = " · ".join(bits)
            refresh_fb()
            def set_fb(rating=None, ftype=None):
                async def _h(e):
                    # toggle off when clicking the same value again
                    if rating is not None:
                        m.rating = None if m.rating == rating else rating
                    if ftype is not None:
                        m.feedback_type = None if m.feedback_type == ftype else ftype
                    refresh_fb()
                    try: save_chat(state["cid"], state["msgs"])
                    except Exception: pass
                    fb_label.update()
                    page.update()
                return _h
            acts.controls += [
                ft.IconButton(ft.Icons.THUMB_UP_OUTLINED, icon_size=15, tooltip=tr("fb_helpful"),
                              on_click=set_fb(rating=5, ftype="helpful")),
                ft.IconButton(ft.Icons.THUMB_DOWN_OUTLINED, icon_size=15, tooltip=tr("fb_bad"),
                              on_click=set_fb(rating=2, ftype="inaccurate")),
                ft.IconButton(ft.Icons.FLAG_OUTLINED, icon_size=15, tooltip=tr("fb_harm"),
                              on_click=set_fb(rating=1, ftype="harmful")),
                fb_label,
            ]
            if m.variants:  # #4 переключатель вариантов
                vi = ft.Text(f'{tr("variants")}: {len(m.variants)+1}', size=11, color=T["muted"])
                async def prev_v(e):
                    if m.variants: m.text, m.variants[-1] = m.variants[-1], m.text
                    md.value = m.text; save_chat(state["cid"], state["msgs"]); page.update()
                acts.controls += [ft.IconButton(ft.Icons.ARROW_LEFT, icon_size=15, on_click=prev_v), vi]
        wrap.controls.append(acts)
        chat_box.controls.append(wrap)
        if state.get("stick", True): scroll_end()
        return md

    def upd_tokens():  # view: токены + прогресс (математика — views.token_stats)
        try: cl = int(settings.get("context_length", 8192) or 8192)
        except (ValueError, TypeError): cl = 8192
        if _token_stats is not None:
            label, frac, _over = _token_stats(state["msgs"], cl, estimate_tokens)
            tok_label.value = f"{label} {tr('tokens')} · {len(state['msgs'])} msg"
        else:
            t = sum(estimate_tokens(m.text or "") for m in state["msgs"])
            frac = min(1.0, t / cl) if cl else 0
            tok_label.value = f"~{t} / {cl} {tr('tokens')} · {len(state['msgs'])} msg"
        try:
            ctx_bar.value = frac
            ctx_bar.color = "#EF5350" if frac >= 0.9 else (th["accent"])
            ctx_bar.update()
        except Exception: pass
        try: page.update()
        except Exception: pass

    def render_all():
        """Перерисовать чат с учётом фильтров (поиск / только оценённые)."""
        chat_box.controls.clear()
        if _visible_messages is not None:
            msgs = _visible_messages(state["msgs"], state.get("chat_filter", ""),
                                     state.get("rated_only", False))
        else:
            q = (state.get("chat_filter") or "").lower()
            rated_only = state.get("rated_only", False)
            msgs = [m for m in state["msgs"]
                    if not (rated_only and not (m.rating or m.feedback_type))
                    and not (q and q not in (m.text or "").lower())]
        for m in msgs:
            add_bubble(m)
        page.update()

    async def load_models(e=None):
        try: models = await client.fetch_models()
        except RuntimeError as ex: _log.warning("fetch_models failed: %s", ex); show_e(str(ex)); set_conn(False); return
        if not models: models = [DEFAULT_MODEL]
        model_dd.options = [ft.DropdownOption(k, k) for k in models]
        saved = settings.get("model") or ""
        model_dd.value = saved if saved in models else models[0]
        try: model_dd.update()
        except Exception: pass
        state["models_n"] = len(models)
        set_conn(True, tr("models_n", n=len(models))); page.update()
        await ensure_model_loaded()

    async def ensure_model_loaded():
        """Загрузить выбранную модель на сервере, выгрузить предыдущую."""
        sel = model_dd.value or DEFAULT_MODEL
        if sel == state.get("loaded_model"): return
        prev = state.get("loaded_model")
        status.value = tr("model_loading", m=sel); page.update()
        if prev and prev != sel:
            try:
                await client.unload_model(prev)  # сначала выгрузить текущую…
                _log.info("unloaded previous model: %s", prev)
            except Exception as ex:
                _log.warning("unload_model(%s) failed: %s", prev, ex)
            state["loaded_model"] = None
        try:
            await client.load_model(sel)  # …только потом грузить новую
        except Exception as ex:
            _log.error("load_model(%s) failed: %s", sel, ex)
            show_e(tr("e_load_model", m=sel, e=ex))
            return
        state["loaded_model"] = sel
        try:
            if store is not None: store._emit("model:loaded")
        except Exception: pass
        status.value = tr("model_loaded", m=sel); page.update()

    async def on_model_change(e=None):
        state["model_touched"] = True
        persist()  # сразу запомнить выбор
        await ensure_model_loaded()

    # --- автоопределение vision: поведенческая память (settings['no_vision_models']) ---
    def _vision_known_bad(model_id: str) -> bool:
        try: return model_id in (settings.get("no_vision_models") or [])
        except Exception: return False

    def _vision_mark_bad(model_id: str):
        lst = [m for m in (settings.get("no_vision_models") or []) if m != model_id]
        lst.append(model_id)
        settings["no_vision_models"] = lst[-50:]  # кап списка
        save_settings(settings)
        _log.info("model marked as no-vision: %s", model_id)

    def _vision_mark_good(model_id: str):
        lst = [m for m in (settings.get("no_vision_models") or []) if m != model_id]
        if len(lst) != len(settings.get("no_vision_models") or []):
            settings["no_vision_models"] = lst
            save_settings(settings)
            _log.info("model proved vision-capable, unmarked: %s", model_id)

    async def confirm_image_retry() -> bool:
        """Модель без vision отклонила image_url (HTTP 400): повторить без картинок?"""
        fut = asyncio.get_event_loop().create_future()
        def close(v):
            try: page.pop_dialog()
            except Exception: pass
            if not fut.done(): fut.set_result(v)
            page.update()
        page.show_dialog(ft.AlertDialog(
            title=ft.Text(tr("dlg_novision_t")),
            content=ft.Text(tr("dlg_novision_c")),
            actions=[ft.TextButton(tr("retry_noimg"), on_click=lambda e: close(True)),
                     ft.TextButton(tr("cancel"), on_click=lambda e: close(False))],
            actions_alignment=ft.MainAxisAlignment.END))
        return await fut

    async def generate(_strip_images: bool = False):
        _model_id = model_dd.value or DEFAULT_MODEL
        if not settings.get("send_images", True):
            _strip_images = True  # vision выключен в настройках — всегда text-only
        elif not _strip_images and _vision_known_bad(_model_id):
            _strip_images = True  # авто: модель ранее роняла 400 на картинках — без диалога
            _log.info("auto-stripping images for known no-vision model: %s", _model_id)
        try: api = build_api(state["msgs"], effective_system(settings, settings.get("system_prompt", "")),
                             int(settings.get("context_length", 8192) or 0),
                             strip_images=_strip_images)
        except FileError as ex: show_e(str(ex)); return
        try:
            from api_payload import payload_has_images as _has_img, payload_stats as _pstats  # type: ignore
            _has_img_f, _stats_f = _has_img, _pstats
        except ImportError:
            def _has_img_f(a):  # type: ignore
                return any(isinstance(m.get("content"), list) and
                           any(isinstance(p, dict) and p.get("type") == "image_url" for p in m["content"])
                           for m in a)
            def _stats_f(a):  # type: ignore
                return {"messages": len(a), "images": 0}
        st = _stats_f(api)
        _log.info("request model=%s msgs=%s images=%s stripped=%s",
                  model_dd.value or DEFAULT_MODEL, st["messages"], st["images"], _strip_images)
        if _strip_images and st["images"]:
            status.value = (tr("vision_off")
                            if not settings.get("send_images", True)
                            else tr("vision_auto"))
            status.update()
        am = ChatMessage(text="", is_user=False); state["msgs"].append(am)
        md = add_bubble(am)
        # --- визуализация ожидания: спиннер-статус + плейсхолдер в пузыре ---
        md.value = tr("generating")
        status.value = tr("generating"); status.color = th["accent"]; page.update()
        disp, buf, last = "", [], time.time()
        _first_tok = True
        def on_delta(kind, tok):
            nonlocal disp, last, _first_tok
            if _first_tok:
                _first_tok = False
                try:
                    status.value = tr("typing"); page.update()
                except Exception: pass
            buf.append(tok if kind == "content" else f"`{tok}`")
            if time.time() - last > 0.15:  # #9 троттлинг по времени
                disp += "".join(buf); buf.clear(); md.value = disp; md.update()
                if state.get("stick", True): scroll_end()
                last = time.time()
        try:
            content, reasoning = await client.chat_stream(api, model_dd.value or DEFAULT_MODEL, settings, on_delta)
            disp += "".join(buf)
            full = (f"> 💭 {tr('reasoning')}\n{reasoning}\n\n---\n" if reasoning and not content else "") + (content or reasoning or disp)
            am.text = content or reasoning or disp; md.value = am.text
            if not _strip_images and _has_img_f(api):
                _vision_mark_good(_model_id)  # картинки прошли — модель с vision
            save_chat(state["cid"], state["msgs"]); upd_tokens(); status.value = ""; status.color = th["muted"]; page.update()
            scroll_end()
        except StreamCancelled:
            am.text = md.value or ""; save_chat(state["cid"], state["msgs"]); status.value = tr("stopped"); status.color = th["muted"]; page.update()
            scroll_end()
        except RuntimeError as ex:
            _log.error("generate failed: %s", ex)
            state["msgs"].pop(); chat_box.controls.pop()
            if "400" in str(ex) and not _strip_images and _has_img_f(api):
                _log.warning("HTTP 400 with images in payload — offering text-only retry")
                if await confirm_image_retry():
                    _vision_mark_bad(_model_id)  # запомнить: без vision, дальше — автострип
                    await generate(_strip_images=True)
                    return
            show_e(str(ex))

    async def send(e=None):
        if state["sending"]: return
        txt = (inp.value or "").strip()
        if not txt and not state["files"]: return
        persist(); hide_e()
        # контроль переполнения контекста: предложить сжатие
        try: cl = int(settings.get("context_length", 8192) or 8192)
        except (ValueError, TypeError): cl = 8192
        cur = sum(estimate_tokens(m.text or "") for m in state["msgs"]) + estimate_tokens(txt)
        if cur > cl:
            go = await confirm_ctx_overflow(cur, cl)
            if go == "cancel": return
            if go == "compress":
                await summarize()
                cur = sum(estimate_tokens(m.text or "") for m in state["msgs"]) + estimate_tokens(txt)
                if cur > cl: show_e(f"~{cur} / {cl} {tr('tokens')} — {tr('ctx_still_over')}"); return
        state["sending"] = True; set_sending_ui(True); page.update()
        try:
            um = ChatMessage(text=txt, is_user=True)
            for p in state["files"]:
                mime, _ = mimetypes.guess_type(p)
                if mime and mime.startswith("image/"):
                    b64, mt = encode_image(p); um.attachments.append(Attachment(path=p, mime=mt, b64=b64))
                else: um.attachments.append(Attachment(path=p))  # текст извлекается в build_api (#2)
            state["msgs"].append(um); add_bubble(um)
            inp.value = ""; counter.value = f"0 / {S['input_max']}"; state["files"] = []; attach_row.controls.clear()
            # автоназвание чата (#1)
            for c in load_index():
                if c["id"] == state["cid"] and c["title"] in (LANGS["ru"]["me"], LANGS["en"]["me"]):
                    c["title"] = txt[:40]; save_index(load_index()); refresh_sidebar()
            save_chat(state["cid"], state["msgs"]); upd_tokens(); page.update()
            await generate()
        finally: state["sending"] = False; set_sending_ui(False); page.update()

    async def regenerate(e=None):
        if state["sending"]: return
        idx = max((i for i, m in enumerate(state["msgs"]) if m.is_user), default=-1)
        if idx < 0: show_e(tr("e_no_msgs")); return
        old = state["msgs"][idx+1:]
        for o in old:  # сохранить прошлый ответ как вариант (#4)
            if not o.is_user and o.text: pass
        if old and not old[-1].is_user and old[-1].text:
            pass
        del state["msgs"][idx+1:]
        chat_box.controls.clear()
        for m in state["msgs"]: add_bubble(m)
        page.update(); state["sending"] = True; set_sending_ui(True)
        try:
            prev_n = len(state["msgs"])
            await generate()
            if len(state["msgs"]) > prev_n and old and not old[-1].is_user:
                state["msgs"][-1].variants = [o.text for o in old if not o.is_user and o.text][:5] + state["msgs"][-1].variants[:5]
                save_chat(state["cid"], state["msgs"])
        finally: state["sending"] = False; set_sending_ui(False); page.update()

    async def summarize(e=None):  # #3 сжатие
        if len(state["msgs"]) < 4: show_e(tr("e_few_msgs")); return
        keep = state["msgs"][-4:]; old = state["msgs"][:-4]
        txt = "\n".join(f"{'U' if m.is_user else 'A'}: {m.text[:500]}" for m in old)
        try:
            c, _ = await client.chat_stream([{"role": "user", "content": tr("summary_of", t=txt[:8000])}],
                model_dd.value or DEFAULT_MODEL, settings, lambda k, t: None)
            state["msgs"] = [ChatMessage(text=tr("summary_hist", c=c), is_user=False)] + keep
            chat_box.controls.clear()
            for m in state["msgs"]: add_bubble(m)
            save_chat(state["cid"], state["msgs"]); upd_tokens(); page.update()
        except Exception as ex: show_e(str(ex))

    def export(fmt):  # #5
        ms = state["msgs"]
        if fmt == "md":
            out = "\n\n".join(f"**{tr('you') if m.is_user else tr('assistant')}** ({time.strftime('%H:%M', time.localtime(m.ts))}):\n{m.text}" for m in ms)
            (DATA / f"chat_{state['cid']}.md").write_text(out, "utf-8")
        elif fmt == "html":
            body = "".join(f"<p><b>{tr('you') if m.is_user else tr('ai')}:</b> {html_mod.escape(m.text)}</p>" for m in ms)
            (DATA / f"chat_{state['cid']}.html").write_text(f"<html><body>{body}</body></html>", "utf-8")
        elif fmt == "feedback":  # выгрузка оценок в JSONL для fine-tuning (без b64)
            def _nodump(m):
                try: return m.to_dict(include_b64=False)
                except TypeError: return m.to_dict()
            rows = [_nodump(m) for m in ms if not m.is_user and (m.rating or m.feedback_type)]
            if not rows: show_e(tr("e_no_msgs")); return
            p = DATA / f"feedback_{state['cid']}.jsonl"
            p.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows), "utf-8")
        else: save_chat(state["cid"], ms)
        status.value = tr("exported", f=fmt); page.update()

    async def confirm_ctx_overflow(cur: int, cl: int) -> str:
        """Диалог при переполнении контекста: compress / continue / cancel."""
        fut = asyncio.get_event_loop().create_future()
        def close(v):
            try: page.pop_dialog()
            except Exception: pass
            if not fut.done(): fut.set_result(v)
            page.update()
        dlg = ft.AlertDialog(
            title=ft.Text(tr("dlg_ctxfull_t")),
            content=ft.Text(tr("dlg_ctxfull_c", cur=cur, cl=cl)),
            actions=[ft.TextButton(tr("summarize"), on_click=lambda e: close("compress")),
                     ft.TextButton(tr("continue_btn"), on_click=lambda e: close("continue")),
                     ft.TextButton(tr("cancel"), on_click=lambda e: close("cancel"))],
            actions_alignment=ft.MainAxisAlignment.END)
        page.show_dialog(dlg)
        return await fut

    def find_in_chat(e=None):  # полнотекстовый поиск по текущему чату
        q_f = ft.TextField(label=tr("m_find"), autofocus=True,
                           value=state.get("chat_filter", ""))
        def close(e=None):
            try: page.pop_dialog()
            except Exception: pass
            page.update()
        def apply(e=None):
            state["chat_filter"] = (q_f.value or "").strip()
            close(); render_all()
            if state["chat_filter"]:
                n = sum(1 for m in state["msgs"] if state["chat_filter"].lower() in (m.text or "").lower())
                status.value = f"🔍 {n}"; page.update()
        def clear(e=None):
            state["chat_filter"] = ""; close(); render_all()
        page.show_dialog(ft.AlertDialog(title=ft.Text(tr("m_find")),
            content=q_f, actions=[ft.TextButton(tr("clear"), on_click=clear),
            ft.TextButton(tr("cancel"), on_click=close), ft.TextButton(tr("save"), on_click=apply)],
            actions_alignment=ft.MainAxisAlignment.END))

    def toggle_rated_only(e=None):  # фильтр: только оценённые сообщения
        state["rated_only"] = not state.get("rated_only", False)
        render_all()
        n = sum(1 for m in state["msgs"] if m.rating or m.feedback_type)
        status.value = f"★ {n}" if state["rated_only"] else ""; page.update()

    async def pick(e=None):
        fs_ = await picker.pick_files(allow_multiple=True,
            allowed_extensions=["png", "jpg", "jpeg", "webp", "gif", "txt", "py", "md", "pdf", "docx", "xlsx", "csv"])
        if not fs_: return
        for f in fs_:
            if f.path and f.path not in state["files"]: state["files"].append(f.path)
        refresh_attach_row()

    def refresh_attach_row():
        """Превью вложений: миниатюры картинок, чипы файлов, у каждого — крестик удаления."""
        attach_row.controls.clear()
        for p in list(state["files"]):
            mime, _ = mimetypes.guess_type(p)
            try: size = f" {(Path(p).stat().st_size // 1024)} KB"
            except Exception: size = ""
            def _rm(e, path=p):
                try: state["files"].remove(path)
                except ValueError: pass
                refresh_attach_row()
            if mime and mime.startswith("image/"):
                attach_row.controls.append(ft.Row([
                    ft.Container(ft.Image(src=p, width=72, height=56, fit=ft.BoxFit.COVER,
                                          border_radius=8, error_content=ft.Text("🖼")),
                                 tooltip=Path(p).name + size),
                    ft.IconButton(ft.Icons.CLOSE, icon_size=13, tooltip=tr("delete"),
                                  on_click=_rm)],
                    spacing=0, vertical_alignment=ft.CrossAxisAlignment.START))
            else:
                attach_row.controls.append(ft.Chip(
                    label=ft.Text(f"{Path(p).name}{size}"),
                    on_delete=lambda e, path=p: _rm(e, path)))
        page.update()

    def on_inp(e):
        n = len(inp.value or "")
        counter.value = f"{n} / {S['input_max']}"
        counter.color = "#EF5350" if n > S["input_max"] else th["muted"]
        counter.update()
    inp.on_change = on_inp; inp.on_submit = send
    async def load_prompt_file(e=None):  # загрузка System prompt из .txt
        hide_e()
        try:
            res = await picker.pick_files(allow_multiple=False, allowed_extensions=["txt", "md"])
        except Exception as ex: show_e(tr("e_open_dlg", e=ex)); return
        if not res:
            status.value = tr("load_cancel"); page.update(); return
        p = res[0].path if hasattr(res[0], "path") else res[0].get("path")
        try:
            txt = Path(p).read_text(encoding="utf-8-sig")
        except Exception as ex: show_e(tr("e_read_file", e=ex)); return
        sys_f.value = txt; persist(); status.value = tr("prompt_loaded", n=Path(p).name, l=len(txt)); page.update()
    async def save_prompt_file(e=None):  # сохранение System prompt в .txt
        hide_e()
        txt = (sys_f.value or "").strip()
        if not txt: show_e(tr("e_empty_prompt")); return
        try:
            out = await picker.save_file(file_name="system_prompt.txt", allowed_extensions=["txt"])
        except Exception as ex: show_e(tr("e_open_dlg", e=ex)); return
        if not out:
            status.value = tr("save_cancel"); page.update(); return
        p = out.path if hasattr(out, "path") else (out if isinstance(out, str) else None)
        try:
            Path(p).write_text(txt, encoding="utf-8")
        except Exception as ex: show_e(tr("e_save_file", ex=ex)); return
        persist(); status.value = tr("prompt_saved", p=p); page.update()
    search.on_change = lambda e: refresh_sidebar()
    preset_name = ft.TextField(label=tr("preset_name"), hint_text=tr("preset_name_hint"), width=200, dense=True)
    def refresh_presets(sel=None):
        PRESETS.clear(); PRESETS.update(_all_presets())
        preset_dd.options = [ft.DropdownOption(key=k, text=preset_display(k)) for k in PRESETS]
        if sel and sel in PRESETS: preset_dd.value = sel
        elif preset_dd.value not in PRESETS: preset_dd.value = "Обычный"
        try: preset_dd.update()
        except Exception: pass
        page.update()
    def apply_preset(e):
        key = (getattr(getattr(e, "control", None), "value", None)
               or preset_dd.value)  # значение из события надёжнее
        new_text = PRESETS.get(key, "") if key in PRESETS else ""
        sys_f.value = new_text  # очистить + вставить одним присвоением
        try: sys_f.update()
        except Exception: pass
        persist(); page.update()
    def save_preset(e):  # сохранить текущий System prompt как пресет
        hide_e()
        name = (preset_name.value or "").strip()
        txt = (sys_f.value or "").strip()
        if not name: show_e(tr("e_enter_preset")); return
        if not txt: show_e(tr("e_empty_prompt")); return
        if name in BUILTIN_PRESETS: show_e(tr("e_builtin_name")); return
        custom = _load_custom_presets(); custom[name] = txt
        try: _save_custom_presets(custom)
        except Exception as ex: show_e(tr("e_save_preset", ex=ex)); return
        preset_name.value = ""; refresh_presets(sel=name)
        status.value = tr("preset_saved", n=name); page.update()
    def delete_preset(e):  # удалить пользовательский пресет
        name = preset_dd.value
        if name in BUILTIN_PRESETS or name == "Обычный": show_e(tr("e_builtin_del")); return
        custom = _load_custom_presets()
        if name not in custom: show_e(tr("e_no_preset")); return
        del custom[name]
        try: _save_custom_presets(custom)
        except Exception as ex: show_e(tr("e_del_preset", ex=ex)); return
        refresh_presets(sel="Обычный"); apply_preset(None)
        status.value = tr("preset_deleted", n=name); page.update()
    def clear_prompt(e):  # очистить поле System prompt
        sys_f.value = ""; sys_f.update(); persist()
        status.value = tr("prompt_cleared"); page.update()
    # --- Профили связок: модель + пресет + параметры одним кликом ---
    try:
        from repositories import ProfilesRepository as _ProfilesRepository  # type: ignore
        _profiles_repo = _ProfilesRepository(PROFILES_F)
        def _load_profiles() -> dict: return _profiles_repo.load()
        def _save_profiles(p: dict): _profiles_repo.save(p)
    except ImportError:
        def _load_profiles() -> dict:
            try:
                if PROFILES_F.is_file():
                    raw = json.loads(PROFILES_F.read_text("utf-8"))
                    if isinstance(raw, dict): return raw
            except Exception: pass
            return {}
        def _save_profiles(p: dict):
            PROFILES_F.write_text(json.dumps(p, ensure_ascii=False, indent=2), "utf-8")
    profile_dd = ft.Dropdown(label=tr("profile"), width=220)
    profile_name = ft.TextField(label=tr("profile_name"), hint_text=tr("profile_name_hint"), width=200, dense=True)
    def refresh_profiles(sel=None):
        try:
            profs = _load_profiles()
            profile_dd.options = [ft.DropdownOption(key=k, text=k) for k in profs]
            if sel and sel in profs: profile_dd.value = sel
            else: profile_dd.value = None
            profile_dd.update()
        except Exception: pass
    async def apply_profile(e=None):
        name = profile_dd.value
        profs = _load_profiles()
        if not name or name not in profs: return
        p = profs[name]
        hide_e()
        try: tmp.value = float(p.get("temperature", 0.7)); tmp_val.value = f"{float(tmp.value):.2f}"; tmp.update(); tmp_val.update()
        except Exception: pass
        try: topp.value = float(p.get("top_p", 1.0)); topp_val.value = f"{float(topp.value):.2f}"; topp.update(); topp_val.update()
        except Exception: pass
        maxt.value = str(p.get("max_tokens", 2048)); seed.value = str(p.get("seed", -1))
        ctxlen.value = str(p.get("context_length", 8192))
        send_images_cb.value = bool(p.get("send_images", True))
        try:
            for c in (maxt, seed, ctxlen, send_images_cb): c.update()
        except Exception: pass
        if p.get("preset") in PRESETS:
            preset_dd.value = p["preset"]
            try: preset_dd.update()
            except Exception: pass
        if p.get("system_prompt") is not None:
            sys_f.value = p["system_prompt"]
            try: sys_f.update()
            except Exception: pass
        persist(); page.update()
        if p.get("model") and p["model"] != model_dd.value:
            if p["model"] not in [o.key for o in model_dd.options]:
                show_e(tr("e_no_model_on_server", m=p['model'])); return
            model_dd.value = p["model"]
            try: model_dd.update()
            except Exception: pass
            state["model_touched"] = True; persist()
            await ensure_model_loaded()
        status.value = tr("profile_applied", n=name); page.update()
    def save_profile(e):
        hide_e()
        name = (profile_name.value or "").strip()
        if not name: show_e(tr("e_enter_preset")); return
        profs = _load_profiles()
        profs[name] = {"model": model_dd.value or "", "preset": preset_dd.value,
            "system_prompt": sys_f.value or "", "temperature": float(tmp.value or 0.7),
            "top_p": float(topp.value or 1.0), "max_tokens": str(maxt.value or 2048),
            "seed": str(seed.value or -1), "context_length": str(ctxlen.value or 8192),
            "send_images": bool(send_images_cb.value)}
        try: _save_profiles(profs)
        except Exception as ex: show_e(str(ex)); return
        profile_name.value = ""; refresh_profiles(sel=name)
        status.value = tr("profile_saved", n=name); page.update()
    def delete_profile(e):
        name = profile_dd.value
        profs = _load_profiles()
        if not name or name not in profs: show_e(tr("e_no_preset")); return
        del profs[name]
        try: _save_profiles(profs)
        except Exception as ex: show_e(str(ex)); return
        refresh_profiles(); status.value = tr("profile_deleted", n=name); page.update()
    # --- Environment Context: requirements.txt / .env ---
    env_label = ft.Text("", size=12, color=th["muted"])
    def refresh_env_label():
        req = (settings.get("env_requirements") or "")
        nvars = len(settings.get("env_vars") or {})
        bits = []
        if req: bits.append(tr("env_req", n=len(req)))
        if nvars: bits.append(tr("env_vars_n", n=nvars))
        env_label.value = " · ".join(bits) if bits else tr("env_none")
    refresh_env_label()
    def _parse_dotenv(text: str) -> dict:
        out = {}
        for line in (text or "").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line: continue
            k, v = line.split("=", 1)
            k = k.strip()
            if k.startswith("export "): k = k[7:].strip()
            out[k] = "***"  # mask values, keep keys
        return out
    async def load_requirements(e=None):
        hide_e()
        try:
            res = await picker.pick_files(allow_multiple=False, allowed_extensions=["txt"])
        except Exception as ex: show_e(tr("e_open_dlg", e=ex)); return
        if not res: status.value = tr("load_cancel"); page.update(); return
        p = res[0].path if hasattr(res[0], "path") else res[0].get("path")
        try: txt = Path(p).read_text(encoding="utf-8-sig")[:8000]
        except Exception as ex: show_e(tr("e_read_file", e=ex)); return
        settings["env_requirements"] = txt; settings["env_req_name"] = Path(p).name
        save_settings(settings); refresh_env_label(); env_label.update()
        status.value = tr("req_loaded", n=Path(p).name); page.update()
    async def load_dotenv(e=None):
        hide_e()
        try:
            res = await picker.pick_files(allow_multiple=False, allowed_extensions=["env", "txt"])
        except Exception as ex: show_e(tr("e_open_dlg", e=ex)); return
        if not res: status.value = tr("load_cancel"); page.update(); return
        p = res[0].path if hasattr(res[0], "path") else res[0].get("path")
        try: txt = Path(p).read_text(encoding="utf-8-sig")
        except Exception as ex: show_e(tr("e_read_file", e=ex)); return
        settings["env_vars"] = _parse_dotenv(txt)
        save_settings(settings); refresh_env_label(); env_label.update()
        status.value = tr("env_loaded", n=len(settings['env_vars'])); page.update()
    def clear_env(e=None):
        settings["env_requirements"] = ""; settings["env_vars"] = {}
        save_settings(settings); refresh_env_label(); env_label.update(); page.update()
    preset_dd.on_select = apply_preset
    preset_dd.on_blur = apply_preset
    model_dd.on_select = on_model_change
    profile_dd.on_select = apply_profile
    # (7) живые значения слайдеров
    tmp_val = ft.Text(f"{float(tmp.value or 0.7):.2f}", size=12, color=th["atc"], width=36)
    topp_val = ft.Text(f"{float(topp.value or 1.0):.2f}", size=12, color=th["atc"], width=36)
    def on_tmp(e): tmp_val.value = f"{float(tmp.value or 0):.2f}"; tmp_val.update(); persist()
    def on_topp(e): topp_val.value = f"{float(topp.value or 0):.2f}"; topp_val.update(); persist()
    tmp.on_change = on_tmp; topp.on_change = on_topp
    def on_ctx(e=None):
        persist(); upd_tokens()
    try: ctxlen.on_blur = on_ctx
    except Exception: pass
    send_images_cb.on_change = lambda e: persist()
    def switch_theme(e):
        settings["theme"] = "light" if settings.get("theme") == "dark" else "dark"
        save_settings(settings); page.bgcolor = THEMES[settings["theme"]]["bg"]; page.update()
    def font_pm(d):
        settings["font_scale"] = min(1.5, max(0.8, float(settings.get("font_scale", 1)) + d))
        save_settings(settings); status.value = tr("font_n", v=f"{settings['font_scale']:.1f}"); page.update()

    # --- i18n: кнопка языка + живое обновление ---
    def safe_update(*cs):
        for c in cs:
            try: c.update()
            except Exception: pass
    def set_lang(e=None):
        lang = getattr(getattr(e, "control", None), "value", None) or lang_btn.value or "ru"
        if lang not in LANGS: return
        old = CUR["lang"]
        CUR["lang"] = lang
        if _i18n is not None:
            try: _i18n.set_lang(CUR["lang"])
            except ValueError: pass
        settings["lang"] = CUR["lang"]; save_settings(settings)
        # если в поле лежит текст встроенного пресета на старом языке — заменить на новый
        for k, v in BUILTIN_PRESETS.items():
            if isinstance(v, dict) and (sys_f.value or "") == v.get(old, "") and v.get(old):
                sys_f.value = v.get(CUR["lang"], "")
                break
        apply_lang()
    def make_overflow():
        return [
            ft.PopupMenuItem(tr("m_compress"), icon=ft.Icons.COMPRESS_OUTLINED, on_click=summarize),
            ft.PopupMenuItem(tr("m_find"), icon=ft.Icons.SEARCH_OUTLINED, on_click=find_in_chat),
            ft.PopupMenuItem(tr("m_rated"), icon=ft.Icons.STAR_OUTLINE, on_click=toggle_rated_only),
            ft.PopupMenuItem(tr("m_exp_md"), icon=ft.Icons.SHARE_OUTLINED, on_click=lambda e: export("md")),
            ft.PopupMenuItem(tr("m_exp_html"), icon=ft.Icons.SHARE_OUTLINED, on_click=lambda e: export("html")),
            ft.PopupMenuItem(tr("m_exp_json"), icon=ft.Icons.SHARE_OUTLINED, on_click=lambda e: export("json")),
            ft.PopupMenuItem(tr("m_exp_fb"), icon=ft.Icons.FEEDBACK_OUTLINED, on_click=lambda e: export("feedback")),
            ft.PopupMenuItem(tr("m_theme"), icon=ft.Icons.BRIGHTNESS_6_OUTLINED, on_click=switch_theme),
            ft.PopupMenuItem(tr("m_font_up"), icon=ft.Icons.TEXT_INCREASE_OUTLINED, on_click=lambda e: font_pm(0.1)),
            ft.PopupMenuItem(tr("m_font_down"), icon=ft.Icons.TEXT_DECREASE_OUTLINED, on_click=lambda e: font_pm(-0.1))]
    def apply_lang():
        page.title = tr("title")
        UI["side_head"].value = tr("chats")
        search.hint_text = tr("search")
        UI["new_btn"].content = tr("new_chat")
        UI["top_title"].value = tr("title")
        model_dd.label = tr("model")
        UI["refresh"].tooltip = tr("refresh_models")
        if state.get("models_n"):
            dot.bgcolor = "#4CAF50"; conn_t.value = tr("models_n", n=state["models_n"])
        else:
            conn_t.value = state["conn_custom"] if state["conn_custom"] is not None else (
                tr("connected") if state["conn_ok"] else tr("no_conn"))
        lang_btn.label = tr("lang")
        if lang_btn.value != CUR["lang"]:
            lang_btn.value = CUR["lang"]
        UI["overflow"].items = make_overflow()  # пересоздать пункты меню
        UI["overflow"].tooltip = tr("more")
        UI["settings_title"].value = tr("settings")
        UI["sec_model"].value = tr("sec_model"); UI["sec_prompt"].value = tr("sec_prompt")
        preset_dd.label = tr("preset")
        refresh_presets()  # тексты пресетов + подписи на новом языке
        preset_name.label = tr("preset_name"); preset_name.hint_text = tr("preset_name_hint")
        sys_f.label = tr("sys_prompt")
        UI["to_preset"].content = tr("to_preset")
        UI["del_preset"].tooltip = tr("del_preset")
        UI["clear"].content = tr("clear")
        UI["load_txt"].content = tr("load_txt"); UI["save_txt"].content = tr("save_txt")
        UI["attach"].tooltip = tr("attach")
        UI["mic"].tooltip = tr("stt_mic")
        send_images_cb.label = tr("vision_send")
        UI["sec_profiles"].value = tr("sec_profiles")
        UI["sec_env"].value = tr("sec_env")
        profile_dd.label = tr("profile")
        profile_name.label = tr("profile_name"); profile_name.hint_text = tr("profile_name_hint")
        UI["profile_save"].content = tr("to_profile")
        UI["profile_del"].tooltip = tr("del_profile")
        UI["menu"].tooltip = tr("chats")
        refresh_env_label()
        try: env_label.update()
        except Exception: pass
        inp.hint_text = tr("input_hint")
        btn_send.tooltip = tr("send")
        try: btn_stop.tooltip = tr("stop")
        except NameError: pass
        safe_update(UI["side_head"], search, UI["new_btn"], UI["top_title"], model_dd,
                    UI["refresh"], conn_t, lang_btn, UI["overflow"],
                    UI["settings_title"], UI["sec_model"], UI["sec_prompt"],
                    preset_dd, preset_name, sys_f,
                    UI["to_preset"], UI["del_preset"], UI["clear"],
                    UI["load_txt"], UI["save_txt"], UI["attach"], UI["mic"], inp, btn_send,
                    UI["sec_profiles"], UI["sec_env"], profile_dd, profile_name,
                    UI["profile_save"], UI["profile_del"], send_images_cb, env_label)
        try: safe_update(btn_stop)
        except NameError: pass
        try: sys_f.update()
        except Exception: pass
        upd_tokens(); refresh_sidebar()
        chat_box.controls.clear()  # перерисовать подписи пузырей на новом языке
        for m in state["msgs"]: add_bubble(m)
        page.update()
    lang_btn = ft.Dropdown(label=tr("lang"), value=CUR["lang"], width=110,
                           options=[ft.DropdownOption("ru", "RU"), ft.DropdownOption("en", "EN")],
                           on_select=set_lang)

    # --- (4) sidebar + slim-топбар: модель + статус, остальное в «⋮» ---
    UI["side_head"] = ft.Text(tr("chats"), weight=ft.FontWeight.BOLD, color=th["atc"])
    UI["new_btn"] = ft.Button(tr("new_chat"), icon=ft.Icons.ADD_OUTLINED, on_click=lambda e: new_chat())
    sidebar = ft.Container(width=S["sidebar_w"], bgcolor=th["panel"], border_radius=S["radius"], padding=8,
        border=ft.Border.all(1, th["border"]), visible=False,
        content=ft.Column([UI["side_head"], search, chat_list, UI["new_btn"]]))
    def toggle_sidebar(e=None):
        sidebar.visible = not sidebar.visible
        try: sidebar.update()
        except Exception: pass
        page.update()
    UI["menu"] = ft.IconButton(ft.Icons.MENU_OUTLINED, tooltip=tr("chats"), on_click=toggle_sidebar)
    UI["overflow_items"] = make_overflow()
    UI["refresh"] = ft.IconButton(ft.Icons.REFRESH_OUTLINED, tooltip=tr("refresh_models"), on_click=load_models)
    overflow = ft.PopupMenuButton(icon=ft.Icons.MORE_VERT_OUTLINED, tooltip=tr("more"),
        items=UI["overflow_items"])
    UI["overflow"] = overflow
    # --- (7) хелпер секций + кнопки (должны быть созданы ДО топбара: Промпт живёт в нём) ---
    def sec(key, icon, *controls):
        t = ft.Text(tr(key), size=13, weight=ft.FontWeight.BOLD, color=th["atc"])
        UI[key] = t
        return ft.Column([ft.Row([ft.Icon(icon, size=16, color=th["accent"]), t],
                                 spacing=6)] + list(controls), spacing=6)
    UI["settings_title"] = ft.Text(tr("settings"), color=th["atc"])
    UI["to_preset"] = ft.OutlinedButton(tr("to_preset"), icon=ft.Icons.BOOKMARK_ADD_OUTLINED, on_click=save_preset)
    UI["del_preset"] = ft.IconButton(ft.Icons.DELETE_OUTLINE, tooltip=tr("del_preset"), on_click=delete_preset)
    UI["clear"] = ft.OutlinedButton(tr("clear"), icon=ft.Icons.CLEAR_OUTLINED, on_click=clear_prompt)
    UI["load_txt"] = ft.OutlinedButton(tr("load_txt"), icon=ft.Icons.UPLOAD_FILE_OUTLINED, on_click=load_prompt_file)
    UI["save_txt"] = ft.OutlinedButton(tr("save_txt"), icon=ft.Icons.SAVE_OUTLINED, on_click=save_prompt_file)
    UI["attach"] = ft.IconButton(ft.Icons.ATTACH_FILE_OUTLINED, tooltip=tr("attach"), on_click=pick)
    UI["load_req"] = ft.OutlinedButton("requirements.txt", icon=ft.Icons.UPLOAD_FILE_OUTLINED, on_click=load_requirements)
    UI["load_env"] = ft.OutlinedButton(".env", icon=ft.Icons.UPLOAD_FILE_OUTLINED, on_click=load_dotenv)
    UI["clear_env"] = ft.OutlinedButton(tr("clear"), icon=ft.Icons.CLEAR_OUTLINED, on_click=clear_env)
    UI["profile_save"] = ft.OutlinedButton(tr("to_profile"), icon=ft.Icons.BOOKMARK_ADD_OUTLINED, on_click=save_profile)
    UI["profile_del"] = ft.IconButton(ft.Icons.DELETE_OUTLINE, tooltip=tr("del_profile"), on_click=delete_profile)
    UI["sec_profiles"] = ft.Text(tr("sec_profiles"), size=13, weight=ft.FontWeight.BOLD, color=th["atc"])
    UI["top_title"] = ft.Text(tr("title"), weight=ft.FontWeight.BOLD, color=th["atc"])
    # Промпт — в том же ряду, что и модель: топбар из двух колонок (слева модель/статус, справа Промпт)
    preset_dd.width = 170
    sys_f.expand = True
    topbar = ft.Container(bgcolor=th["panel"], border_radius=S["radius"], padding=8,
        border=ft.Border.all(1, th["border"]),
        content=ft.Row([
            ft.Column([
                ft.Row([UI["menu"], ft.Icon(ft.Icons.CHAT_BUBBLE_OUTLINE, color=th["accent"]),
                    UI["top_title"]],
                    spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ft.Row([model_dd, UI["refresh"]], spacing=8,
                       vertical_alignment=ft.CrossAxisAlignment.CENTER),
                ft.Row([dot, conn_t], spacing=4),
                ft.Row([ft.Column([tok_label, ctx_bar], spacing=2), lang_btn, overflow],
                       spacing=8, vertical_alignment=ft.CrossAxisAlignment.CENTER)],
                spacing=6),
            ft.VerticalDivider(width=8, color=th["border"]),
            ft.Column([
                sec("sec_prompt",
                    ft.Icons.CHAT_OUTLINED,
                    ft.Row([preset_dd, sys_f], spacing=8,
                           vertical_alignment=ft.CrossAxisAlignment.START),
                    ft.Row([preset_name, UI["to_preset"], UI["del_preset"], UI["clear"]],
                           wrap=True, spacing=6, run_spacing=4),
                    ft.Row([UI["load_txt"], UI["save_txt"]], wrap=True, spacing=6))],
                spacing=6, expand=True)],
            spacing=8, vertical_alignment=ft.CrossAxisAlignment.START))
    # --- Настройки: Модель + Окружение (Промпт живёт в топбаре) ---
    UI["sec_env"] = ft.Text(tr("sec_env"), size=13, weight=ft.FontWeight.BOLD, color=th["atc"])
    settings_p = ft.ExpansionTile(title=UI["settings_title"],
        leading=ft.Icon(ft.Icons.TUNE_OUTLINED, color=th["accent"]),
        controls=[
        sec("sec_model",
            ft.Icons.SMART_TOY_OUTLINED,
            ft.Row([ft.Text("Temperature", size=13, color=th["atc"], width=90), tmp, tmp_val],
                   vertical_alignment=ft.CrossAxisAlignment.CENTER),
            ft.Row([ft.Text("top_p", size=13, color=th["atc"], width=90), topp, topp_val],
                   vertical_alignment=ft.CrossAxisAlignment.CENTER),
            ft.Row([maxt, seed, ctxlen], wrap=True),
            ft.Row([send_images_cb], wrap=True)),
        ft.Divider(height=4, color=th["border"]),
        ft.Column([ft.Row([ft.Icon(ft.Icons.LAYERS_OUTLINED, size=16, color=th["accent"]),
                           UI["sec_profiles"]],
                          spacing=6),
                   profile_dd,
                   ft.Row([profile_name, UI["profile_save"], UI["profile_del"]], wrap=True)], spacing=6),
        ft.Divider(height=4, color=th["border"]),
        ft.Column([ft.Row([ft.Icon(ft.Icons.DNS_OUTLINED, size=16, color=th["accent"]),
                           UI["sec_env"]],
                          spacing=6),
                   env_label,
                   ft.Row([UI["load_req"], UI["load_env"], UI["clear_env"]], wrap=True)], spacing=6)])
    # --- (5) плавающий док ввода ---
    btn_send = ft.IconButton(ft.Icons.ARROW_UPWARD_ROUNDED, tooltip=tr("send"),
        on_click=send, style=ft.ButtonStyle(bgcolor=th["accent"], color="white", shape=ft.CircleBorder()))
    def set_sending_ui(v: bool):
        """Переключить док в режим генерации: показать Стоп, заблокировать отправку."""
        try:
            btn_stop.visible = v
            btn_send.disabled = v
            btn_stop.update(); btn_send.update()
        except Exception: pass
    def do_stop(e=None):
        client.cancel()
        status.value = tr("stopped"); page.update()
    async def do_listen(e=None):
        try:
            from voice import listen  # type: ignore
        except ImportError:
            show_e(tr("e_no_stt")); return
        mic = UI["mic"]
        # --- визуально: красная активная кнопка ---
        _old = (mic.icon, mic.bgcolor, mic.icon_color, mic.tooltip)
        try:
            mic.icon = ft.Icons.MIC_ROUNDED
            mic.bgcolor = "#EF5350"
            mic.icon_color = "white"
            mic.tooltip = tr("mic_stop_title")
            mic.update()
        except Exception: pass
        status.value = tr("mic_listening"); status.color = "#EF5350"; page.update()
        try:
            text = await asyncio.get_running_loop().run_in_executor(
                None, listen, "ru-RU" if CUR["lang"] == "ru" else "en-US")
            inp.value = ((inp.value or "") + " " + text).strip()
            inp.update(); on_inp(None)
        except Exception as ex: show_e(str(ex))
        finally:  # --- возврат к базовому виду ---
            try:
                mic.icon, mic.bgcolor, mic.icon_color, mic.tooltip = _old
                mic.update()
            except Exception: pass
            status.value = ""; status.color = th["muted"]; page.update()
    UI["mic"] = ft.IconButton(ft.Icons.MIC_OUTLINED, tooltip=tr("stt_mic"), on_click=do_listen)
    btn_stop = ft.IconButton(ft.Icons.STOP_CIRCLE_OUTLINED, tooltip=tr("stop"), on_click=do_stop,
        visible=False, style=ft.ButtonStyle(color="#EF5350"))
    dock = ft.Container(bgcolor=th["input_bg"], border_radius=S["radius"] + 4, padding=8,
        border=ft.Border.all(1, th["border"]),
        shadow=ft.BoxShadow(blur_radius=12, color=th["shadow"], offset=ft.Offset(0, -2)),
        content=ft.Column(spacing=4, controls=[
            attach_row,
            ft.Row([UI["attach"],
                    inp, UI["mic"], btn_stop, btn_send],
                   vertical_alignment=ft.CrossAxisAlignment.END),
            ft.Row([status, ft.Container(expand=True), counter],
                   vertical_alignment=ft.CrossAxisAlignment.CENTER)]))
    chat_wrap = ft.Container(content=chat_box, expand=True)
    page.add(ft.Row([sidebar,
                     ft.Column([topbar, settings_p, err, chat_wrap, dock],
                               expand=True, spacing=S["gap"])],
                    expand=True, vertical_alignment=ft.CrossAxisAlignment.START))
    def on_keyboard(e):
        k = (e.key or "").lower()
        if e.ctrl and k == "enter":
            asyncio.create_task(send())
        elif e.ctrl and k == "k":
            new_chat()
        elif e.ctrl and k == "f":
            find_in_chat()
        elif k == "escape" and state.get("sending"):
            do_stop()
    page.on_keyboard_event = on_keyboard
    def stop_lm_server():
        """Остановить сервер LM Studio через CLI (best effort, только если lms доступен)."""
        import shutil, subprocess
        lms = shutil.which("lms")
        if not lms:
            _log.info("lms CLI not found — server left running")
            return
        try:
            r = subprocess.run([lms, "server", "stop"], timeout=15, capture_output=True, text=True)
            _log.info("lms server stop: %s", (r.stdout or r.stderr or "ok").strip()[:200])
        except Exception as ex:
            _log.warning("lms server stop failed: %s", ex)

    def unload_via_cli(model_id: str | None = None):
        """Fallback: выгрузка через `lms unload`, если HTTP API не сработал."""
        import shutil, subprocess
        lms = shutil.which("lms")
        if not lms:
            return
        # сначала точечно модель, потом всё остальное
        cmds = []
        if model_id:
            cmds.append([lms, "unload", model_id])
        cmds.append([lms, "unload", "--all"])
        for cmd in cmds:
            try:
                r = subprocess.run(cmd, timeout=30, capture_output=True, text=True)
                _log.info("lms %s: %s", " ".join(cmd[1:]), (r.stdout or r.stderr or "ok").strip()[:200])
                if r.returncode == 0:
                    break
            except Exception as ex:
                _log.warning("lms unload failed: %s", ex)

    async def _exit_network():
        """Выгрузка модели + стоп сервера. Вызывается под shield — переживает отмену задачи."""
        try:
            lm = state.get("loaded_model") or (model_dd.value or None)
            if lm:
                try:
                    await asyncio.wait_for(client.unload_model(lm), timeout=15)
                    _log.info("unloaded model on exit: %s", lm)
                except Exception as ex:
                    _log.warning("unload on exit via API failed: %s — пробую lms unload", ex)
                    try:
                        await asyncio.get_running_loop().run_in_executor(None, unload_via_cli, lm)
                    except BaseException as ex2:
                        _log.warning("unload via CLI failed: %s", ex2)
                state["loaded_model"] = None
        except BaseException as ex:
            _log.warning("exit unload block failed: %s", ex)
        # стоп сервера — СТРОГО после выгрузки, пока клиент ещё жив
        try:
            await asyncio.get_running_loop().run_in_executor(None, stop_lm_server)
        except BaseException as ex:
            _log.warning("server stop on exit failed: %s", ex)
        try: await client.close()
        except BaseException: pass

    async def on_app_close(e=None):
        _log.info("app closing")
        try:  # 1) геометрия окна — быстро и синхронно, первым делом
            try: settings["window_maximized"] = bool(page.window.maximized)
            except Exception: pass
            if not settings.get("window_maximized"):
                if page.window.width: settings["window_width"] = int(page.window.width)
                if page.window.height: settings["window_height"] = int(page.window.height)
                if page.window.left is not None: settings["window_left"] = int(page.window.left)
                if page.window.top is not None: settings["window_top"] = int(page.window.top)
            _log.info("saved geometry: %sx%s max=%s", settings.get("window_width"),
                      settings.get("window_height"), settings.get("window_maximized"))
        except Exception as ex:
            _log.warning("save window geometry failed: %s", ex)
        try:  # 2) остальные настройки
            persist()
        except Exception as ex:
            _log.warning("persist on exit failed: %s", ex)
        try:  # 3) сеть под shield — flet отменяет задачу при закрытии окна
            await asyncio.shield(_exit_network())
        except BaseException as ex:
            _log.warning("exit network cancelled: %s", type(ex).__name__)
        _log.info("exit handler done")
    page.on_close = on_app_close

    # шаг 6: flet-обновления едут от стора через подписчика (views.ViewBinder)
    if store is not None and _ViewBinder is not None:
        try: _ViewBinder(store, on_tokens=upd_tokens).bind()
        except Exception: _log.exception("view binder failed")

    idx = load_index()
    if not idx: new_chat()
    else: open_chat(idx[0]["id"])
    refresh_sidebar(); refresh_profiles(); await load_models()

if __name__ == "__main__":
    ft.run(main)
